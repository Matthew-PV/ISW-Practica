import json

import openpyxl


path = r"C:\Users\mattp\AppData\Local\Temp\planb-cs30-readonly.xlsx"
workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)

result = {"cs30": [], "matthew_tasks": []}

sheet = workbook["CS-30"]
for row_number in range(1, 45):
    values = [sheet.cell(row_number, column).value for column in range(1, 7)]
    if any(value is not None for value in values):
        result["cs30"].append([row_number, *values])

for sheet_name in workbook.sheetnames:
    if not sheet_name.startswith("CS-"):
        continue
    sheet = workbook[sheet_name]
    for row_number in range(17, 60):
        description = sheet.cell(row_number, 2).value
        owner = sheet.cell(row_number, 4).value
        estimate = sheet.cell(row_number, 5).value
        if isinstance(owner, str) and owner.strip().casefold() == "matthew":
            result["matthew_tasks"].append(
                {
                    "story": sheet_name,
                    "row": row_number,
                    "description": description,
                    "estimate": estimate,
                }
            )

print(json.dumps(result, ensure_ascii=True, default=str, indent=2))
